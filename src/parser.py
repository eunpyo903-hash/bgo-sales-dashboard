"""
비고컴퍼니 매출현황 - Excel Parser

책임:
  1) data/ 폴더에서 최신 월(YYMM)의 지점별 Excel 파일을 자동으로 찾는다.
  2) 각 지점 Excel 파일에서 표준 데이터 항목을 추출한다.

셀 매핑은 실제 지점 Excel 7개(역삼/선정릉/대치/송파/장안/가락/삼전, 2609)를
openpyxl로 직접 열어 Sheet 구성(매출/매출순위/페이롤), 병합 셀, 수식과 수식
결과를 비교 분석한 뒤 확정했다. 7개 파일 모두 동일한 셀 위치를 사용한다.

[매출] 시트 레이아웃 (공통 확인됨)
  J1        = 월 (예: "9월")
  U1        = 지점명 (예: "역삼점")
  C3/E3/G3  = 목표 FC / 목표 PT / 목표 TOTAL
  C4/E4/G4  = 실적 FC / 실적 PT / 실적 TOTAL   (수식 결과)
  C5/E5/G5  = 달성률 FC / PT / TOTAL           (0~1 소수, 수식 결과)
  K8/L8     = FC 신규 건수 / 신규 매출
  M8/N8     = FC 재등록 건수 / 재등록 매출
  R8/S8     = PT 신규 건수 / 신규 매출
  T8/U8     = PT 재등록 건수 / 재등록 매출
  M11       = 총 카드매출 (card)
  O11       = 카드 외 매출 (cash)
  Q11       = 계좌 매출 (account)
  S11       = 소득공제 (deduction)
  U11       = 실 카드 외 매출 (real_cash) = O11 + Q11 - S11
  B15:B(끝) = 상세 거래 내역의 날짜(일자) → 최신 거래일(기준일) 계산에 사용

[매출순위] 시트 레이아웃 (공통 확인됨, PT 개인매출 순위 원본)
  C열 = 직급 (매니저/팀장/시니어/팀원)
  D열 = PT 트레이너 성명 (비어 있으면 미배정 슬롯 → 제외)
  E열 = 해당 트레이너 PT 개인 매출 (수식 결과)
  F열 = 지점명
"""

import os
from collections import defaultdict

from openpyxl import load_workbook

from .utils import BRANCHES, parse_filename

SHEET_SALES = "매출"
SHEET_PT_RANKING = "매출순위"

# [매출] 시트 셀 주소
CELL_TARGET_FC = "C3"
CELL_TARGET_PT = "E3"
CELL_TARGET_TOTAL = "G3"
CELL_ACTUAL_FC = "C4"
CELL_ACTUAL_PT = "E4"
CELL_ACTUAL_TOTAL = "G4"
CELL_RATE_FC = "C5"
CELL_RATE_PT = "E5"
CELL_RATE_TOTAL = "G5"

CELL_FC_NEW_COUNT = "K8"
CELL_FC_NEW_SALES = "L8"
CELL_FC_RENEW_COUNT = "M8"
CELL_FC_RENEW_SALES = "N8"
CELL_PT_NEW_COUNT = "R8"
CELL_PT_NEW_SALES = "S8"
CELL_PT_RENEW_COUNT = "T8"
CELL_PT_RENEW_SALES = "U8"

CELL_CARD = "M11"
CELL_CASH = "O11"
CELL_ACCOUNT = "Q11"
CELL_DEDUCTION = "S11"
CELL_REAL_CASH = "U11"

DETAIL_START_ROW = 15
DETAIL_DATE_COL = 2  # B열

# [매출순위] 시트 컬럼
PT_RANK_START_ROW = 4
PT_RANK_TITLE_COL = 3   # C열
PT_RANK_NAME_COL = 4    # D열
PT_RANK_SALES_COL = 5   # E열
PT_RANK_BRANCH_COL = 6  # F열

REQUIRED_CELLS = [
    CELL_TARGET_FC, CELL_TARGET_PT, CELL_TARGET_TOTAL,
    CELL_ACTUAL_FC, CELL_ACTUAL_PT, CELL_ACTUAL_TOTAL,
    CELL_RATE_FC, CELL_RATE_PT, CELL_RATE_TOTAL,
    CELL_FC_NEW_COUNT, CELL_FC_NEW_SALES, CELL_FC_RENEW_COUNT, CELL_FC_RENEW_SALES,
    CELL_PT_NEW_COUNT, CELL_PT_NEW_SALES, CELL_PT_RENEW_COUNT, CELL_PT_RENEW_SALES,
]


class DuplicateFileWarning(Exception):
    """중복 파일 경고를 나타내기 위한 용도 (raise 하지 않고 메시지 수집용으로도 사용 가능)"""
    pass


class ParserNotReadyError(Exception):
    """실제 Excel 구조와 다른(예상 Sheet/셀이 없는) 경우 발생시키는 예외"""
    pass


def scan_data_dir(data_dir: str):
    """
    data_dir 안의 모든 xlsx 파일명을 정규식으로 분석하여
    {yymm: {branch_no: [file_info, ...]}} 형태로 반환한다.
    """
    found = defaultdict(lambda: defaultdict(list))

    if not os.path.isdir(data_dir):
        return found

    for fname in os.listdir(data_dir):
        if not fname.lower().endswith((".xlsx", ".xlsm")):
            continue
        if fname.startswith("~$"):  # Excel 임시 잠금 파일 제외
            continue

        parsed = parse_filename(fname)
        if not parsed:
            continue

        full_path = os.path.join(data_dir, fname)
        found[parsed["yymm"]][parsed["branch_no"]].append({
            "filename": fname,
            "path": full_path,
            "branch_name_in_file": parsed["branch_name_in_file"],
            "mtime": os.path.getmtime(full_path),
        })

    return found


def select_latest_files(data_dir: str, warnings: list = None):
    """
    가장 최신 YYMM을 자동 선택하고, 지점별로 파일 하나씩 확정한다.

    같은 월/같은 지점에 파일이 여러 개면:
      1차: 파일명 정확도 (표준 패턴에 더 정확히 맞는 파일 우선)
      2차: modified time (최신 수정 파일 우선)
    을 기준으로 선택하고, 중복이 있었다는 WARNING을 warnings 리스트에 남긴다.

    반환: (period_yymm, {branch_no: file_info}) 또는 (None, {}) - 파일이 전혀 없을 때
    """
    if warnings is None:
        warnings = []

    scanned = scan_data_dir(data_dir)
    if not scanned:
        return None, {}

    latest_yymm = sorted(scanned.keys())[-1]
    return latest_yymm, _pick_branch_files(latest_yymm, scanned[latest_yymm], warnings)


def select_files_for_month(data_dir: str, yymm: str, warnings: list = None) -> dict:
    """지정한 YYMM의 지점별 파일을 하나씩 확정한다. 해당 월 파일이 없으면 {}."""
    if warnings is None:
        warnings = []
    scanned = scan_data_dir(data_dir)
    if yymm not in scanned:
        return {}
    return _pick_branch_files(yymm, scanned[yymm], warnings)


def _pick_branch_files(yymm: str, branch_files: dict, warnings: list) -> dict:
    selected = {}
    for branch_no, candidates in branch_files.items():
        if len(candidates) == 1:
            selected[branch_no] = candidates[0]
            continue

        # 1차: 표준 파일명 패턴과의 정확도 우선 정렬, 2차: mtime 최신
        def sort_key(c):
            standard_name = f"{yymm}_b_{branch_no}_{BRANCHES[branch_no]}점_총매출"
            exact_match = 1 if c["filename"].startswith(standard_name) else 0
            return (exact_match, c["mtime"])

        candidates_sorted = sorted(candidates, key=sort_key, reverse=True)
        selected[branch_no] = candidates_sorted[0]

        other_names = ", ".join(c["filename"] for c in candidates_sorted[1:])
        warnings.append(
            f"WARNING: {yymm} {BRANCHES[branch_no]}점 파일이 {len(candidates)}개 발견되어 "
            f"'{candidates_sorted[0]['filename']}' 파일을 선택했습니다. (제외: {other_names})"
        )

    return selected


def _require_sheet(wb, sheet_name, file_path):
    if sheet_name not in wb.sheetnames:
        raise ParserNotReadyError(
            f"'{os.path.basename(file_path)}'에서 '{sheet_name}' Sheet를 찾을 수 없습니다. "
            f"(발견된 Sheet: {', '.join(wb.sheetnames)})"
        )
    return wb[sheet_name]


def _last_transaction_day(ws) -> int:
    """[매출] 시트 상세 거래 내역(B열, 날짜)에서 가장 최근 일자를 찾는다."""
    last_day = None
    for row in ws.iter_rows(min_row=DETAIL_START_ROW, min_col=DETAIL_DATE_COL,
                             max_col=DETAIL_DATE_COL):
        value = row[0].value
        if isinstance(value, (int, float)):
            last_day = int(value) if last_day is None else max(last_day, int(value))
    return last_day


NEW_RENEW_CELLS = {
    "fc_new_count": CELL_FC_NEW_COUNT, "fc_new_sales": CELL_FC_NEW_SALES,
    "fc_renew_count": CELL_FC_RENEW_COUNT, "fc_renew_sales": CELL_FC_RENEW_SALES,
    "pt_new_count": CELL_PT_NEW_COUNT, "pt_new_sales": CELL_PT_NEW_SALES,
    "pt_renew_count": CELL_PT_RENEW_COUNT, "pt_renew_sales": CELL_PT_RENEW_SALES,
}


def _is_num(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _recompute_new_renew(ws) -> dict:
    """
    상세 거래 내역(15행~)에서 신규/재등록 집계를 직접 다시 센다.
    Excel의 K8:U8 수식(SUBTOTAL, 필터로 숨겨진 행 제외)과 같은 규칙이다.
      FC 건수: G열이 신규/재등록인 행 중 D열(회원권 기간)이 숫자인 행 수
      PT 건수: G열이 신규/재등록인 행 중 E열(PT 세션)이 숫자인 행 수
      FC 매출: I,J,K열 합 / PT 매출: O,P,Q열 합
    """
    result = {}
    for key, label in (("new", "신규"), ("renew", "재등록")):
        fc_count = fc_sales = pt_count = pt_sales = 0
        for r in range(DETAIL_START_ROW, 515):
            if ws.row_dimensions[r].hidden or ws.cell(row=r, column=7).value != label:
                continue
            if _is_num(ws.cell(row=r, column=4).value):
                fc_count += 1
            if _is_num(ws.cell(row=r, column=5).value):
                pt_count += 1
            fc_sales += sum(ws.cell(row=r, column=c).value for c in (9, 10, 11)
                            if _is_num(ws.cell(row=r, column=c).value))
            pt_sales += sum(ws.cell(row=r, column=c).value for c in (15, 16, 17)
                            if _is_num(ws.cell(row=r, column=c).value))
        result[f"fc_{key}_count"] = fc_count
        result[f"fc_{key}_sales"] = int(fc_sales) if float(fc_sales).is_integer() else fc_sales
        result[f"pt_{key}_count"] = pt_count
        result[f"pt_{key}_sales"] = int(pt_sales) if float(pt_sales).is_integer() else pt_sales
    return result


def extract_branch_data(file_path: str, branch_no: int) -> dict:
    """지점 Excel 1개의 [매출] Sheet에서 표준 데이터 항목을 추출한다."""
    wb = load_workbook(filename=file_path, data_only=True)
    ws = _require_sheet(wb, SHEET_SALES, file_path)

    missing = [addr for addr in REQUIRED_CELLS if ws[addr].value is None]
    if missing:
        raise ParserNotReadyError(
            f"'{os.path.basename(file_path)}'의 [매출] Sheet에서 다음 셀 값을 찾을 수 없습니다: "
            f"{', '.join(missing)}"
        )

    def pct(cell_addr):
        """Excel 소수(0.483...) 형태의 달성률을 퍼센트 숫자로 변환한다."""
        return round(ws[cell_addr].value * 100, 1)

    recomputed = _recompute_new_renew(ws)
    new_renew = {}
    notes = []
    for key, addr in NEW_RENEW_CELLS.items():
        saved = ws[addr].value
        if saved != recomputed[key]:
            notes.append(
                f"WARNING: {BRANCHES[branch_no]}점 {key}: Excel 저장값({saved})과 상세내역 재계산값"
                f"({recomputed[key]})이 달라 재계산값을 사용했습니다. ({addr})"
            )
        new_renew[key] = recomputed[key]

    return {
        **new_renew,
        "parser_notes": notes,
        "target_fc": ws[CELL_TARGET_FC].value,
        "target_pt": ws[CELL_TARGET_PT].value,
        "target_total": ws[CELL_TARGET_TOTAL].value,

        "actual_fc": ws[CELL_ACTUAL_FC].value,
        "actual_pt": ws[CELL_ACTUAL_PT].value,
        "actual_total": ws[CELL_ACTUAL_TOTAL].value,

        "rate_fc": pct(CELL_RATE_FC),
        "rate_pt": pct(CELL_RATE_PT),
        "rate_total": pct(CELL_RATE_TOTAL),

        "card": ws[CELL_CARD].value,
        "cash": ws[CELL_CASH].value,
        "account": ws[CELL_ACCOUNT].value,
        "deduction": ws[CELL_DEDUCTION].value,
        "real_cash": ws[CELL_REAL_CASH].value,

        "last_transaction_day": _last_transaction_day(ws),
    }


def extract_pt_ranking(file_path: str, branch_no: int) -> list:
    """
    [매출순위] Sheet에서 트레이너별 PT 개인매출을 추출한다.
    D열(성명)이 비어 있는 슬롯(미배정)은 제외한다.
    """
    wb = load_workbook(filename=file_path, data_only=True)
    ws = _require_sheet(wb, SHEET_PT_RANKING, file_path)

    branch_name = BRANCHES[branch_no]
    rows = []
    for r in range(PT_RANK_START_ROW, ws.max_row + 1):
        name = ws.cell(row=r, column=PT_RANK_NAME_COL).value
        if name is None or str(name).strip() == "":
            continue
        sales = ws.cell(row=r, column=PT_RANK_SALES_COL).value
        branch_cell = ws.cell(row=r, column=PT_RANK_BRANCH_COL).value
        title = ws.cell(row=r, column=PT_RANK_TITLE_COL).value

        rows.append({
            "name": str(name).strip(),
            "branch_no": branch_no,
            "branch_name": (branch_cell or f"{branch_name}점").replace("점", "") or branch_name,
            "sales": sales if isinstance(sales, (int, float)) else 0,
            "title": str(title).strip() if title else None,
        })

    return rows


WALKIN_KEYWORD = "워크인"
DETAIL_ETC_COL = 6       # F열 (기타)
DETAIL_MANAGER_COL = 8   # H열 (담당자)
DETAIL_MEMO_COL = 22     # V열 (비고)
# [매출순위] E열 수식과 같은 PT 매출 컬럼: O,P,Q(PT매출) + T(PT 환불)
DETAIL_PT_SALES_COLS = (15, 16, 17, 20)


def extract_walkin_pt_sales(file_path: str) -> dict:
    """
    [매출] 상세 거래 내역 중 비고(V열) 또는 기타(F열)에 '워크인'이 적힌 행의
    PT 매출을 담당자(H열)별로 합산한다. 반환: {담당자: 금액}
    """
    wb = load_workbook(filename=file_path, data_only=True)
    ws = _require_sheet(wb, SHEET_SALES, file_path)

    result = defaultdict(float)
    for r in range(DETAIL_START_ROW, 515):
        memo = f"{ws.cell(row=r, column=DETAIL_MEMO_COL).value or ''} {ws.cell(row=r, column=DETAIL_ETC_COL).value or ''}"
        if WALKIN_KEYWORD not in memo.replace(" ", ""):
            continue
        manager = ws.cell(row=r, column=DETAIL_MANAGER_COL).value
        if not manager:
            continue
        result[str(manager).strip()] += sum(
            ws.cell(row=r, column=c).value for c in DETAIL_PT_SALES_COLS
            if _is_num(ws.cell(row=r, column=c).value)
        )
    return {k: int(v) if float(v).is_integer() else v for k, v in result.items()}


def open_workbook(file_path: str, data_only: bool = True):
    """수식 결과값을 읽을 때는 data_only=True를 기본으로 사용한다."""
    return load_workbook(filename=file_path, data_only=data_only)
